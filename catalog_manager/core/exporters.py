"""
core/exporters.py — 결과 파일 내보내기 (교정 카탈로그, 유사매칭, URL오류, CSV, 새 카탈로그)
"""
import csv
import os
import re

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

from core.utils import (
    normalize_name, normalize_fee, extract_date_range, fuzzy_ratio,
    get_lines, norm_day, norm_time, clean_instructor, dedup,
)
from config import BASE_URL, STORE_CODE


# ── 공통 스타일 ───────────────────────────────────────────────────────────

_BLUE_FONT  = Font(name="Malgun Gothic", color="0000CD", underline="single")
_BLUE_FILL  = PatternFill("solid", start_color="DCE8FF", end_color="DCE8FF")
_ORANGE_FONT = Font(name="Malgun Gothic", color="B8670A", underline="single")
_ORANGE_FILL = PatternFill("solid", start_color="FFF3DC", end_color="FFF3DC")
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_AUTHOR      = "교정프로그램"


def _mark_cell(cell, font, fill, comment_text: str):
    """셀에 색상 서식과 코멘트를 적용."""
    cell.font = font
    cell.fill = fill
    try:
        if cell.comment is None:
            cell.comment = Comment(comment_text, _AUTHOR)
        else:
            cell.comment.text += "\n" + comment_text
    except Exception:
        pass


def _get_col_map(ws) -> dict[str, int]:
    """워크시트에서 헤더 행을 찾아 {헤더명: 열번호} 딕셔너리 반환."""
    col_map = {}
    for row_idx in range(1, 11):
        row_vals = [str(c.value or "").replace("\n", "").strip() for c in ws[row_idx]]
        if any("강좌명" in v for v in row_vals):
            for cell in ws[row_idx]:
                v = str(cell.value or "").replace("\n", "").strip()
                if "강좌명" in v and "소" not in v:   col_map["강좌명"] = cell.column
                elif v == "요일":                      col_map["요일"] = cell.column
                elif "강좌일자" in v:                  col_map["강좌일자"] = cell.column
                elif "강좌시간" in v:                  col_map["강좌시간"] = cell.column
                elif "강좌비" in v:                    col_map["강좌비"] = cell.column
                elif "강사명" in v:                    col_map["강사명"] = cell.column
                elif "랜딩" in v:                      col_map["랜딩 주소"] = cell.column
            break
    return col_map


# ── 교정 카탈로그 내보내기 ────────────────────────────────────────────────

_FIELD_TO_COL = {
    "요일":   "요일",
    "강좌시간": "강좌시간",
    "강좌일자": "강좌일자",
    "강좌비":  "강좌비",
    "재료비":  "강좌비",
    "강사명":  "강사명",
}


def export_corrected_catalog(
    catalog_path: str,
    comparison_results: list,
    url_results: list,
    output_path: str,
):
    """불일치 셀에 파란색 밑줄을 적용한 교정 카탈로그 Excel 저장."""
    wb = load_workbook(catalog_path)

    # 비교 결과 마킹
    for res in comparison_results:
        sn = res["sheet"]
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        col_map = _get_col_map(ws)
        erow = int(res["cat_row"]["_excel_row"])

        for (field, cat_val, plan_val) in res["mismatches"]:
            cols = []
            if field in _FIELD_TO_COL:
                cols.append(_FIELD_TO_COL[field])
            elif field == "갯수 불일치":
                cols.extend(["요일", "강좌시간", "강좌일자", "랜딩 주소"])
            elif field == "휴강서식 오류":
                cols.append("강좌일자")

            hint = str(plan_val)[:60] + ("..." if len(str(plan_val)) > 60 else "")
            for ck in cols:
                if ck in col_map:
                    try:
                        _mark_cell(ws.cell(row=erow, column=col_map[ck]),
                                   _BLUE_FONT, _BLUE_FILL, f"[{field}] 기획자료: {hint}")
                    except Exception:
                        pass

        # 유사 매칭 → 강좌명 셀 오렌지 표시
        if res["match_type"] == "fuzzy" and "강좌명" in col_map:
            try:
                c = ws.cell(row=erow, column=col_map["강좌명"])
                plan_name = str(res["plan_row"]["course_name"]) if res["plan_row"] is not None else "-"
                _mark_cell(c, _ORANGE_FONT, _ORANGE_FILL, f"[기획자료 강좌명]: {plan_name}")
            except Exception:
                pass

    # URL 오류 마킹
    url_issues_map: dict = {}
    for ur in url_results:
        if ur["status"] != "ok":
            key = (ur["sheet"], ur["course"][:30])
            url_issues_map.setdefault(key, []).extend(ur["issues"])

    for res in comparison_results:
        sn = res["sheet"]
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        col_map = _get_col_map(ws)
        ckey = (sn, str(res["cat_row"].get("course_name", "")).replace("\n", " ")[:30])
        if ckey in url_issues_map and "랜딩 주소" in col_map:
            erow = int(res["cat_row"]["_excel_row"])
            try:
                _mark_cell(
                    ws.cell(row=erow, column=col_map["랜딩 주소"]),
                    _BLUE_FONT, _BLUE_FILL,
                    "\n".join(url_issues_map[ckey]),
                )
            except Exception:
                pass

    wb.save(output_path)


# ── 유사강좌명 검토 파일 ──────────────────────────────────────────────────

def export_fuzzy_review(comparison_results: list, output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "강좌명 유사매칭 검토"

    headers = ["시트", "카탈로그 강좌명", "기획자료 강좌명", "유사도", "요일", "시간", "강사", "기타불일치"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1E2235", end_color="1E2235")

    for res in comparison_results:
        if res["match_type"] != "fuzzy":
            continue
        plan = res["plan_row"]
        chk  = lambda f: "✓" if not any(m[0] == f for m in res["mismatches"]) else "✗"  # noqa
        ws.append([
            res["sheet"],
            str(res["cat_row"].get("course_name", "")).replace("\n", " ")[:60],
            str(plan["course_name"]).replace("\n", " ")[:60] if plan is not None else "-",
            f"{res['fuzzy_score']:.0%}",
            chk("요일"), chk("강좌시간"), chk("강사명"),
            ", ".join(m[0] for m in res["mismatches"] if m[0] not in ("요일", "강좌시간", "강사명")),
        ])

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max(len(str(c.value or "")) for c in col) + 4, 55
        )
    wb.save(output_path)


# ── URL 오류 목록 ─────────────────────────────────────────────────────────

def export_url_errors(url_results: list, output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "URL오류목록"
    ws.append(["시트", "강좌명", "lectCode", "상태", "이슈"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1E2235", end_color="1E2235")
    for r in url_results:
        if r["status"] != "ok":
            ws.append([r["sheet"], r["course"], r.get("lect_code", ""),
                       r["status"], " | ".join(r["issues"])])
    wb.save(output_path)


# ── 비교결과 CSV ──────────────────────────────────────────────────────────

def export_comparison_csv(comparison_results: list, output_path: str):
    rows = []
    for res in comparison_results:
        mm_str = "; ".join(f"{f}({c}→{p})" for f, c, p in res["mismatches"])
        rows.append({
            "시트":          res["sheet"],
            "카탈로그 강좌명": str(res["cat_row"].get("course_name", "")).replace("\n", " ")[:60],
            "매칭유형":       res["match_type"],
            "유사도":         f"{res['fuzzy_score']:.0%}",
            "불일치항목":     mm_str,
        })
    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else [])
        writer.writeheader()
        writer.writerows(rows)


# ── 새 카탈로그 생성 ──────────────────────────────────────────────────────

def export_catalog_from_scratch(
    landing_df: pd.DataFrame,
    template_path: str,
    year: str,
    smst: str,
    output_path: str,
    default_dates: dict,
    planning_df: pd.DataFrame | None = None,
):
    """
    카탈로그_공란 양식에 랜딩 URL을 채워 저장합니다.

    URL 매칭 기준:
      - 템플릿(카탈로그_공란)의 강좌명을 랜딩주소 파일의 강좌명과 비교하여 lectCode를 찾습니다.
      - 기획자료(planning_df)는 선택 사항입니다.
        * 제공되면: 요일·강좌시간·강좌일자·강사명·강좌비도 함께 채웁니다.
        * 없으면:   랜딩 URL 열만 채웁니다.
    """
    wb = load_workbook(template_path)

    def clean_for_match(s):
        s = re.sub(r"\[.*?]", "", str(s))
        s = re.sub(r"\(.*?\)", "", s)
        return re.sub(r"\s+", "", s).strip()

    def get_base_name(name):
        n = normalize_name(name)
        n = re.sub(r"\s*\[.*?]\s*$", "", n)
        n = re.sub(r"\s*\(.*?\)\s*$", "", n)
        return n.strip()

    def get_suffix(name):
        n = re.sub(r"^\[.*?]\s*", "", str(name)).strip()
        m = re.search(r"\[(.*?)]$|\((.*?)\)$", n)
        return (m.group(1) or m.group(2) or "").strip() if m else ""

    # ── 랜딩주소 레코드 인덱싱 ──────────────────────────────────────────
    landing_records = [
        {
            "clean_name": clean_for_match(r.get("course_name", "")),
            "day":  re.sub(r"[^가-힣]", "", str(r.get("day", ""))),
            "time": re.sub(r"\D", "", str(r.get("time", ""))),
            "code": str(r.get("lect_code", "")).strip(),
        }
        for _, r in landing_df.iterrows()
    ]

    def find_lect_code(cat_name: str, day: str = "", time_str: str = "") -> str:
        """
        템플릿 강좌명을 기준으로 랜딩주소 파일에서 lectCode를 찾습니다.
        요일·시간이 주어지면 매칭 정밀도를 높입니다.
        """
        q = clean_for_match(cat_name)
        best_code, best_score = None, -1
        clean_d = re.sub(r"[^가-힣]", "", day)
        t_norm  = re.sub(r"\D", "", time_str)

        for lr in landing_records:
            # 요일이 다르면 제외
            if clean_d and lr["day"] and clean_d != lr["day"]:
                continue
            sc = fuzzy_ratio(q, lr["clean_name"])
            # 시간까지 일치하면 보너스
            if t_norm and lr["time"] and t_norm == lr["time"]:
                sc += 0.5
            if sc > best_score and sc > 0.6:
                best_score, best_code = sc, lr["code"]

        if best_code:
            return (
                f"{BASE_URL}?yearCode={year}&smstCode={smst}"
                f"&storeCode={STORE_CODE}&lectCode={best_code}"
            )
        return "랜딩코드 없음"

    # ── 기획자료가 있을 때 — 나머지 필드도 채우기 위한 그루핑 ──────────────
    plan_groups: dict[str, list] = {}
    ordered_keys: list[str] = []
    if planning_df is not None:
        for _, row in planning_df.iterrows():
            base = get_base_name(str(row["course_name"]))
            if base not in plan_groups:
                plan_groups[base] = []
                ordered_keys.append(base)
            plan_groups[base].append(row)

    def build_row_data_from_plan(prows, cat_name: str) -> dict:
        """기획자료 행 목록으로 전체 필드를 구성하고, URL은 cat_name 기준으로 조회."""
        days, times, dates, insts, fees, urls = [], [], [], [], [], []

        for pr in prows:
            p_days  = get_lines(pr.get("day"))
            p_times = get_lines(pr.get("time"))
            p_insts = get_lines(pr.get("instructor"))
            suffix  = get_suffix(str(pr["course_name"]))
            dr      = extract_date_range(str(pr.get("course_name", "")))

            for i in range(max(len(p_days), len(p_times), 1)):
                d_val = p_days[i] if i < len(p_days) else (p_days[0] if p_days else "")
                t_val = p_times[i] if i < len(p_times) else (p_times[0] if p_times else "")
                i_val = p_insts[i] if i < len(p_insts) else (p_insts[0] if p_insts else "")

                clean_d = re.sub(r"[^가-힣]", "", str(d_val))
                day_str = f"[{clean_d}]" + (f" {suffix}" if suffix else "")
                days.append(day_str)
                times.append(t_val)
                insts.append(i_val)
                dates.append(dr if dr else default_dates.get(clean_d, ""))

                # URL: 기획자료 강좌명 아닌 카탈로그(템플릿) 강좌명 기준으로 조회
                urls.append(find_lect_code(cat_name, clean_d, str(t_val)))

            cf  = normalize_fee(pr.get("fee"))
            mat = pr.get("material")
            f_str = f"{cf:,}원" if cf else ""
            if pd.notna(mat) and str(mat).strip() not in ("nan", "", "NaN", "None"):
                f_str = (f_str + "\n" if f_str else "") + f"재료비 {str(mat).strip()}"
            if f_str:
                fees.append(f_str)

        return {
            "day":        "\n".join(str(d).replace("~", "-") for d in days if str(d).strip()),
            "time":       "\n".join(str(t).strip().replace("~", "-") for t in times if str(t).strip()),
            "date":       "\n".join(str(d).strip().replace("~", "-") for d in dates if str(d).strip()),
            "instructor": "\n".join(dedup([str(i).strip() for i in insts if str(i).strip()])),
            "fee":        "\n".join(dedup(fees)),
            "url":        "\n".join(urls),
        }

    def _find_header(ws):
        for ri in range(1, 11):
            vals = [str(c.value or "").replace("\n", "").strip() for c in ws[ri]]
            if any("강좌명" in v for v in vals):
                return ri
        return 1

    def _col_map(ws, header_row):
        return {
            str(c.value or "").replace("\n", "").strip(): i + 1
            for i, c in enumerate(ws[header_row]) if c.value
        }

    def _read_cell_day_time(ws, row_idx, cm):
        """해당 행의 요일·시간 셀 값을 읽어 반환 (URL 매칭 정밀도 향상용)."""
        day_val  = ""
        time_val = ""
        for h, ci in cm.items():
            v = str(ws.cell(row=row_idx, column=ci).value or "")
            if h == "요일":
                day_val = v.split("\n")[0]  # 첫 번째 요일만
            elif "강좌시간" in h:
                time_val = v.split("\n")[0]
        return day_val, time_val

    for sname in wb.sheetnames:
        ws = wb[sname]
        hr = _find_header(ws)
        cm = _col_map(ws, hr)
        if "강좌명" not in cm:
            continue

        has_data = any(
            ws.cell(row=r, column=cm["강좌명"]).value
            for r in range(hr + 1, min(hr + 5, ws.max_row + 1))
        )

        if has_data:
            # ── MODE 1: 템플릿에 강좌명이 이미 채워진 경우 ──────────────
            for row_idx in range(hr + 1, ws.max_row + 1):
                cat_name = ws.cell(row=row_idx, column=cm["강좌명"]).value
                if not cat_name:
                    continue

                cat_name_str = str(cat_name)
                base         = get_base_name(cat_name_str)

                if plan_groups:
                    # 기획자료 있으면 전체 필드 채움
                    prows = plan_groups.get(base)
                    if not prows:
                        best, bkey = 0, None
                        for pk in plan_groups:
                            sc = fuzzy_ratio(base, pk)
                            if sc > best:
                                best, bkey = sc, pk
                        if best > 0.65:
                            prows = plan_groups[bkey]
                    if prows:
                        _fill_row(ws, row_idx, cm, base,
                                  build_row_data_from_plan(prows, cat_name_str))
                        continue

                # 기획자료 없거나 매칭 실패 → URL만 채움
                day_val, time_val = _read_cell_day_time(ws, row_idx, cm)
                url_col = next((ci for h, ci in cm.items() if "랜딩" in h), None)
                if url_col:
                    cell = ws.cell(row=row_idx, column=url_col)
                    cell.value = find_lect_code(cat_name_str, day_val, time_val)
                    cell.alignment = _CENTER_WRAP

        else:
            # ── MODE 2: 완전히 빈 양식 ──────────────────────────────────
            # 기획자료가 없으면 빈 양식에는 채울 수 없음
            if not plan_groups:
                continue

            cur = hr + 1
            for base in ordered_keys:
                data = build_row_data_from_plan(plan_groups[base], base)
                _fill_row(ws, cur, cm, base, data)
                lines = max(
                    data[k].count("\n") + 1
                    for k in ["day", "time", "date", "instructor", "url", "fee"]
                )
                ws.row_dimensions[cur].height = lines * 16.5
                cur += 1

    wb.save(output_path)


def _fill_row(ws, row_idx: int, cm: dict, base_name: str, data: dict):
    """워크시트의 지정 행에 데이터를 채웁니다."""
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    field_map = {
        "강좌명":  base_name,
        "요일":    data["day"],
        "강좌시간": data["time"],
        "강좌일자": data["date"],
        "강사명":   data["instructor"],
        "랜딩 주소": data["url"],
    }
    for h_key, col_idx in cm.items():
        val = None
        for pattern, v in field_map.items():
            if pattern in h_key and ("소" not in h_key if pattern == "강좌명" else True):
                val = v
                break
        if "강좌비" in h_key or "수강료" in h_key:
            val = data["fee"]
        if val is not None:
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.alignment = align
